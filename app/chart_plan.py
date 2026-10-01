"""
Plan comptable double référentiel : US GAAP <-> SYCEBNL.

Principe : chaque compte du grand livre porte
  - un code et un intitulé US GAAP        (colonnes number / name)
  - un compte et un intitulé SYCEBNL      (colonnes syc_code / syc_name)
  - une section US GAAP (gaap_section)    -> sert à bâtir le rapport US GAAP
  - un coefficient (par défaut 1)         -> appliqué dans le rapport SYCEBNL
Les écritures sont saisies une seule fois ; les deux versions des rapports
sont calculées à partir des mêmes écritures, en regroupant les comptes
selon ces correspondances. Tout est modifiable depuis « Plan comptable ».
"""
from .chart_data import STANDARD_CHART
from .database import ACCOUNT_TYPES

# section -> (libellé affiché, sens normal : "D" débit / "C" crédit)
GAAP_SECTIONS = {
    "CA":   ("Actif courant (Current assets)", "D"),
    "NCA":  ("Actif non courant (Non-current assets)", "D"),
    "CL":   ("Passif courant (Current liabilities)", "C"),
    "LTL":  ("Passif long terme (Long-term liabilities)", "C"),
    "EQ":   ("Capitaux propres (Equity)", "C"),
    "REV":  ("Produits (Revenue)", "C"),
    "COGS": ("Coût des ventes (COGS)", "D"),
    "OPEX": ("Charges d'exploitation (Operating expenses)", "D"),
    "OINC": ("Autres produits (Other income)", "C"),
    "OEXP": ("Autres charges (Other expenses)", "D"),
    "TAX":  ("Impôt sur le résultat (Income tax)", "D"),
}

ASSET_SECTIONS = ("CA", "NCA")
LIABILITY_SECTIONS = ("CL", "LTL", "EQ")
INCOME_SECTIONS = ("REV", "OINC")            # produits (crédit)
EXPENSE_SECTIONS = ("COGS", "OPEX", "OEXP", "TAX")  # charges (débit)

REVIEW_STATUSES = ["", "Inchangé", "Modifié", "À confirmer", "Sans équivalent"]

# section -> type fonctionnel par défaut (utilisé par les listes déroulantes des autres écrans)
DEFAULT_TYPE_FOR_SECTION = {
    "CA": "AUTRE_ACTIF_CT", "NCA": "ACTIF_LONG_TERME",
    "CL": "PASSIF_COURT_TERME", "LTL": "PASSIF_LONG_TERME", "EQ": "CAPITAUX_PROPRES",
    "REV": "PRODUITS", "COGS": "COUT_DES_VENTES", "OPEX": "DEPENSES",
    "OINC": "AUTRES_PRODUITS", "OEXP": "DEPENSES", "TAX": "DEPENSES",
}

# type fonctionnel -> section, pour les anciens comptes sans section
_SECTION_FROM_TYPE = {
    "TRESORERIE": "CA", "COMPTES_CLIENTS": "CA", "STOCK": "CA", "AUTRE_ACTIF_CT": "CA",
    "ACTIF_LONG_TERME": "NCA",
    "COMPTES_FOURNISSEURS": "CL", "CARTE_CREDIT": "CL", "PASSIF_COURT_TERME": "CL",
    "PASSIF_LONG_TERME": "LTL", "CAPITAUX_PROPRES": "EQ",
    "PRODUITS": "REV", "COUT_DES_VENTES": "COGS", "DEPENSES": "OPEX",
    "AUTRES_PRODUITS": "OINC", "AUTRES_DEPENSES": "OEXP",
}


def has_syc(code):
    """Vrai si le compte a une vraie correspondance SYCEBNL."""
    c = (code or "").strip().upper()
    return c not in ("", "N/A", "NA", "-")


def effective_section(acc):
    """Section US GAAP du compte ; déduite du type fonctionnel pour les anciens comptes."""
    sec = acc["gaap_section"] if "gaap_section" in acc.keys() else None
    if sec in GAAP_SECTIONS:
        return sec
    return _SECTION_FROM_TYPE.get(acc["account_type"], "OPEX")


def guess_section(syc_code):
    """Section US GAAP probable d'après un compte SYCEBNL (à valider par l'utilisateur)."""
    c = (syc_code or "").strip()
    if not has_syc(c):
        return "CA"
    if c[0] == "2":
        return "NCA"
    if c[0] == "3":
        return "CA"
    if c[0] == "4":
        return "CL" if c[:2] in ("40", "42", "43", "44", "45", "46") else "CA"
    if c[0] == "5":
        return "CL" if c.startswith("56") else "CA"
    if c[0] == "1":
        if c[:2] in ("18", "19"):
            return "LTL"
        return "EQ"
    if c[0] == "6":
        if c[:3] in ("601", "602", "603"):
            return "COGS"
        return "OEXP" if c.startswith("67") else "OPEX"
    if c[0] == "7":
        return "OINC" if c[:2] in ("77", "78") else "REV"
    if c[0] == "8":
        return "OINC" if c[:2] in ("82", "84", "86", "88") else "OEXP"
    return "CA"


# ---------------------------------------------------------------------- #
# Classement SYCEBNL (pour le rapport version SYCEBNL)
# ---------------------------------------------------------------------- #
BS_ASSET_GROUPS = ["Actif immobilisé", "Actif circulant", "Trésorerie-Actif"]
BS_LIAB_GROUPS = [
    "Fonds propres et ressources assimilées",
    "Dettes financières et ressources assimilées",
    "Passif circulant",
    "Trésorerie-Passif",
]


def syc_balance_group(acc):
    """Groupe du bilan SYCEBNL dans lequel tombe le compte."""
    sec = effective_section(acc)
    code = (acc["syc_code"] or "").strip() if "syc_code" in acc.keys() else ""
    if sec in ASSET_SECTIONS:
        if has_syc(code):
            if code[0] == "2":
                return BS_ASSET_GROUPS[0]
            if code[0] == "5":
                return BS_ASSET_GROUPS[2]
            return BS_ASSET_GROUPS[1]
        return BS_ASSET_GROUPS[0] if sec == "NCA" else BS_ASSET_GROUPS[1]
    # passif / fonds propres
    if has_syc(code):
        if code[0] == "5":
            return BS_LIAB_GROUPS[3]
        if code[0] == "4":
            return BS_LIAB_GROUPS[2]
        if code[:2] in ("18", "19"):
            return BS_LIAB_GROUPS[1]
        if code[0] == "1":
            return BS_LIAB_GROUPS[0]
    return {"EQ": BS_LIAB_GROUPS[0], "LTL": BS_LIAB_GROUPS[1]}.get(sec, BS_LIAB_GROUPS[2])


def syc_result_group(acc):
    """'exploitation' | 'financier' | 'hao' pour un compte de résultat."""
    sec = effective_section(acc)
    code = (acc["syc_code"] or "").strip() if "syc_code" in acc.keys() else ""
    if has_syc(code):
        if code[0] == "8":
            return "hao"
        if code[:2] in ("67", "77"):
            return "financier"
        return "exploitation"
    return "financier" if sec in ("OINC", "OEXP") else "exploitation"


# ---------------------------------------------------------------------- #
# Chargement / import / export
# ---------------------------------------------------------------------- #
_COLS = ("number, name, syc_code, syc_name, mapping_type, coefficient, review_status, "
         "remark, gaap_section, account_type, currency, is_active")


def _insert_account(db, d):
    db.conn.execute(
        f"INSERT INTO accounts ({_COLS}, detail_type) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (d["number"], d["name"], d.get("syc_code", ""), d.get("syc_name", ""),
         d.get("mapping_type", ""), d.get("coefficient", 1.0), d.get("review_status", ""),
         d.get("remark", ""), d["gaap_section"], d["account_type"],
         d.get("currency", "GHS"), 1 if d.get("is_active", True) else 0,
         ACCOUNT_TYPES.get(d["account_type"], ("", ""))[1]),
    )


def load_standard_chart(db):
    """Ajoute les comptes du plan standard absents de la base (par code). Retourne le nombre ajouté."""
    existing = {r["number"] for r in db.query("SELECT number FROM accounts")}
    added = 0
    for d in STANDARD_CHART:
        if d["number"] in existing:
            continue
        _insert_account(db, d)
        added += 1
    db.conn.commit()
    return added


EXPORT_HEADERS = [
    "Code US GAAP (type)", "Intitulé du compte US GAAP", "Compte SYCEBNL indicatif",
    "Intitulé SYCEBNL", "Type de correspondance / remarque", "Coefficient", "Statut de revue",
    "Ancien compte SYSCOHADA", "Justification / action",
    "Section US GAAP", "Type fonctionnel", "Devise", "Actif (1/0)",
]


def export_chart_xlsx(db, path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Correspondance"
    ws.append(EXPORT_HEADERS)
    for a in db.query("SELECT * FROM accounts ORDER BY number, name"):
        ws.append([
            a["number"] or "", a["name"], a["syc_code"] or "", a["syc_name"] or "",
            a["mapping_type"] or "", a["coefficient"] if a["coefficient"] is not None else 1,
            a["review_status"] or "", "", a["remark"] or "",
            a["gaap_section"] or effective_section(a), a["account_type"], a["currency"],
            a["is_active"],
        ])
    for col, w in zip("ABCDEFGHIJKLM", (14, 36, 14, 44, 20, 11, 14, 14, 50, 14, 22, 9, 10)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def import_chart_xlsx(db, path):
    """
    Importe un fichier de correspondance (même structure que Correspondance_US_GAAP_SYCEBNL.xlsx
    ou qu'un export de l'application). Mise à jour par code US GAAP ; création si absent.
    Retourne (créés, mis_à_jour).
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Correspondance"] if "Correspondance" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return 0, 0

    def cell(r, i):
        v = r[i] if i < len(r) else None
        return "" if v is None else str(v).strip()

    created = updated = 0
    for r in rows[1:]:
        number, name = cell(r, 0), cell(r, 1)
        if not number or not name:
            continue
        syc_code, syc_name = cell(r, 2), cell(r, 3)
        try:
            coef = float(cell(r, 5).replace(",", ".")) if cell(r, 5) else 1.0
        except ValueError:
            coef = 1.0
        std = next((s for s in STANDARD_CHART if s["number"] == number), None)
        section = cell(r, 9).upper()
        if section not in GAAP_SECTIONS:
            section = std["gaap_section"] if std else guess_section(syc_code)
        atype = cell(r, 10)
        if atype not in ACCOUNT_TYPES:
            atype = std["account_type"] if std else DEFAULT_TYPE_FOR_SECTION[section]
        data = {
            "number": number, "name": name, "syc_code": syc_code, "syc_name": syc_name,
            "mapping_type": cell(r, 4), "coefficient": coef, "review_status": cell(r, 6),
            "remark": cell(r, 8), "gaap_section": section, "account_type": atype,
            "currency": cell(r, 11) or "GHS",
            "is_active": cell(r, 12) not in ("0", "False", "false"),
        }
        found = db.query("SELECT id FROM accounts WHERE number=?", (number,))
        if found:
            db.conn.execute(
                """UPDATE accounts SET name=?, syc_code=?, syc_name=?, mapping_type=?,
                          coefficient=?, review_status=?, remark=?, gaap_section=?,
                          account_type=?, currency=?, is_active=? WHERE id=?""",
                (data["name"], data["syc_code"], data["syc_name"], data["mapping_type"],
                 data["coefficient"], data["review_status"], data["remark"],
                 data["gaap_section"], data["account_type"], data["currency"],
                 1 if data["is_active"] else 0, found[0]["id"]),
            )
            updated += 1
        else:
            _insert_account(db, data)
            created += 1
    db.conn.commit()
    return created, updated
