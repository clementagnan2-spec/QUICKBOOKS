"""
Import Excel : écritures comptables, balance N-1 (à-nouveaux) et balance N (clôture).

Chaque import suit deux temps :
  1. prepare_*()  : lit le fichier, contrôle tout, ne modifie RIEN dans la base ;
                    retourne un Plan (résumé, erreurs, fonction apply).
  2. plan.apply() : écrit dans la base, en une seule transaction (tout ou rien).

Comptes reconnus dans le fichier, dans cet ordre :
  code US GAAP  ->  compte SYCEBNL (s'il ne correspond qu'à un seul compte)  ->  intitulé exact.
"""
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Callable, List

from .chart_plan import effective_section, INCOME_SECTIONS, EXPENSE_SECTIONS

TOL = 0.01
DOC_ENTRIES = "IMPORT"
DOC_OPENING = "BALANCE_N-1"
DOC_CLOSING = "BALANCE_N"


class ImportDataError(Exception):
    pass


@dataclass
class Plan:
    summary: str = ""
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    apply: Callable = None          # appelée sans argument, retourne un message de fin
    needs_replace: bool = False     # N-1 : une balance N-1 existe déjà et sera remplacée


# ---------------------------------------------------------------------- #
# Lecture / conversions
# ---------------------------------------------------------------------- #
def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join("".join(ch if ch.isalnum() else " " for ch in s.lower()).split())


def _clean_code(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\u00a0", "").replace(" ", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    x = float(s)
    return -x if neg else x


def _date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        from openpyxl.utils.datetime import from_excel
        return from_excel(v).date().isoformat()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"date illisible « {s} »")


def _read_rows(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _find_header(rows, required):
    """Cherche (index de ligne, {rôle: index de colonne}) ; required = rôles obligatoires."""
    for i, row in enumerate(rows[:20]):
        cols = {}
        for j, cell in enumerate(row):
            h = _norm(cell)
            if not h:
                continue
            if "date" in h and "date" not in cols:
                cols["date"] = j
            elif ("piece" in h or h in ("n", "no", "num", "numero", "reference", "ref", "n piece",
                                        "numero piece", "doc", "document")) and "piece" not in cols:
                cols["piece"] = j
            elif ("libelle" in h or "description" in h or "memo" in h or "narration" in h) \
                    and "label" not in cols:
                cols["label"] = j
            elif h in ("compte", "account", "code", "code compte", "n compte", "numero compte",
                       "numero de compte", "code us gaap", "compte us gaap", "compte sycebnl",
                       "code sycebnl", "code du compte") and "account" not in cols:
                cols["account"] = j
            elif ("intitule" in h or h in ("nom", "nom du compte", "account name", "name")) \
                    and "accname" not in cols:
                cols["accname"] = j
            elif "debit" in h and "debit" not in cols:
                cols["debit"] = j
            elif "credit" in h and "credit" not in cols:
                cols["credit"] = j
            elif ("solde" in h or "balance" in h) and "solde" not in cols:
                cols["solde"] = j
            elif h in ("journal", "type", "code journal") and "journal" not in cols:
                cols["journal"] = j
        if all(any(r in cols for r in group) if isinstance(group, tuple) else group in cols
               for group in required):
            return i, cols
    raise ImportDataError(
        "En-têtes introuvables. Utilisez le bouton « Modèle Excel » pour obtenir le bon format.")


class _Resolver:
    def __init__(self, db):
        self.by_number, self.by_syc, self.by_name, self.by_id = {}, {}, {}, {}
        for a in db.query("SELECT * FROM accounts WHERE is_active=1"):
            self.by_id[a["id"]] = a
            if a["number"]:
                self.by_number[a["number"].strip()] = a
            if (a["syc_code"] or "").strip().upper() not in ("", "N/A"):
                self.by_syc.setdefault(a["syc_code"].strip(), []).append(a)
            self.by_name[_norm(a["name"])] = a

    def find(self, code, name=""):
        """Retourne (compte, None) ou (None, message d'erreur)."""
        code = _clean_code(code)
        if code in self.by_number:
            return self.by_number[code], None
        cands = self.by_syc.get(code, [])
        if len(cands) == 1:
            return cands[0], None
        if len(cands) > 1:
            nums = ", ".join(c["number"] or c["name"] for c in cands)
            return None, (f"compte SYCEBNL {code} ambigu (correspond à : {nums}) — "
                          "utilisez le code US GAAP")
        a = self.by_name.get(_norm(name)) or self.by_name.get(_norm(code))
        if a:
            return a, None
        return None, f"compte « {code or name} » introuvable dans le plan comptable"


def _fmt(x):
    return f"{x:,.2f}".replace(",", " ")


def _errors_head(errors, n=15):
    shown = errors[:n]
    if len(errors) > n:
        shown.append(f"… et {len(errors) - n} autre(s) erreur(s)")
    return shown


# ---------------------------------------------------------------------- #
# 1. Écritures comptables
# ---------------------------------------------------------------------- #
def prepare_entries(db, path):
    rows = _read_rows(path)
    hdr_i, c = _find_header(rows, ["date", "account", ("debit", "credit")])
    res = _Resolver(db)
    errors, warnings = [], []
    lines = []  # dict(line, date, piece, label, journal, acc_id, debit, credit)

    for n, row in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
        def g(k):
            return row[c[k]] if k in c and c[k] < len(row) else None
        if all(v in (None, "") for v in row):
            continue
        try:
            debit, credit = _num(g("debit")), _num(g("credit"))
        except ValueError:
            errors.append(f"Ligne {n} : montant illisible")
            continue
        if abs(debit) < 1e-9 and abs(credit) < 1e-9:
            continue
        if debit < 0 or credit < 0:
            errors.append(f"Ligne {n} : montant négatif (utilisez la colonne opposée)")
            continue
        try:
            d = _date(g("date"))
        except ValueError as e:
            errors.append(f"Ligne {n} : {e}")
            continue
        acc, err = res.find(g("account"), g("accname"))
        if err:
            errors.append(f"Ligne {n} : {err}")
            continue
        lines.append({"n": n, "date": d, "piece": _clean_code(g("piece")),
                      "label": str(g("label") or "").strip(),
                      "journal": str(g("journal") or "").strip(),
                      "acc": acc, "debit": debit, "credit": credit})

    # regroupement en écritures
    groups, keyed = [], {}
    open_group = None
    for l in lines:
        if l["piece"]:
            k = (l["date"], l["piece"])
            if k not in keyed:
                keyed[k] = []
                groups.append(keyed[k])
            keyed[k].append(l)
        else:  # sans n° de pièce : lignes consécutives jusqu'à l'équilibre
            if open_group is None or open_group[0]["date"] != l["date"]:
                open_group = []
                groups.append(open_group)
            open_group.append(l)
            if abs(sum(x["debit"] - x["credit"] for x in open_group)) < TOL:
                open_group = None

    for g_ in groups:
        diff = round(sum(x["debit"] - x["credit"] for x in g_), 2)
        if abs(diff) >= TOL:
            ref = g_[0]["piece"] or f"du {g_[0]['date']}"
            errors.append(f"Écriture {ref} (ligne {g_[0]['n']}) déséquilibrée : écart {_fmt(diff)}")

    # doublons
    to_post, dup = [], 0
    for g_ in groups:
        piece, d = g_[0]["piece"], g_[0]["date"]
        label = g_[0]["label"]
        total = round(sum(x["debit"] for x in g_), 2)
        key_col, key_val = ("doc_number", piece) if piece else ("memo", label)
        if key_val and db.query(
            f"""SELECT 1 FROM journal_entries je WHERE je.date=? AND je.{key_col}=? AND
                ROUND((SELECT SUM(debit) FROM journal_lines WHERE entry_id=je.id),2)=?""",
                (d, key_val, total)):
            dup += 1
            continue
        to_post.append(g_)
    if dup:
        warnings.append(f"{dup} écriture(s) déjà présente(s) (même date, même n° de pièce ou libellé, même montant) : ignorée(s).")

    tot_d = round(sum(x["debit"] for g_ in to_post for x in g_), 2)
    summary = (f"{len(to_post)} écriture(s) à importer ({sum(len(g_) for g_ in to_post)} lignes)\n"
               f"Total débit = total crédit = {_fmt(tot_d)}")

    def apply():
        cn = db.conn
        try:
            for g_ in to_post:
                first = g_[0]
                cur = cn.execute(
                    "INSERT INTO journal_entries (date, doc_type, doc_number, memo) VALUES (?,?,?,?)",
                    (first["date"], first["journal"] or DOC_ENTRIES, first["piece"], first["label"]))
                for x in g_:
                    cn.execute(
                        "INSERT INTO journal_lines (entry_id, account_id, description, debit, credit) "
                        "VALUES (?,?,?,?,?)",
                        (cur.lastrowid, x["acc"]["id"], x["label"], x["debit"], x["credit"]))
            cn.commit()
        except Exception:
            cn.rollback()
            raise
        return f"{len(to_post)} écriture(s) importée(s)."

    return Plan(summary, _errors_head(errors) if errors else [], warnings,
                apply if not errors and to_post else None)


# ---------------------------------------------------------------------- #
# Lecture d'une balance (commun N-1 / N)
# ---------------------------------------------------------------------- #
def _read_balance(db, path):
    """Retourne ({account_id: (compte, solde_net_débit)}, erreurs)."""
    rows = _read_rows(path)
    hdr_i, c = _find_header(rows, ["account", ("debit", "solde")])
    res = _Resolver(db)
    out, errors = {}, []
    for n, row in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
        def g(k):
            return row[c[k]] if k in c and c[k] < len(row) else None
        code = _clean_code(g("account"))
        name = str(g("accname") or "")
        if not code and not name:
            continue
        if _norm(code).startswith("total") or _norm(name).startswith("total"):
            continue
        try:
            if "debit" in c or "credit" in c:
                nd = _num(g("debit")) - _num(g("credit"))
            else:
                nd = _num(g("solde"))
        except ValueError:
            errors.append(f"Ligne {n} : montant illisible")
            continue
        acc, err = res.find(code, name)
        if err:
            if abs(nd) >= 1e-9:
                errors.append(f"Ligne {n} : {err}")
            continue
        prev = out.get(acc["id"], (acc, 0.0))[1]
        out[acc["id"]] = (acc, round(prev + nd, 2))
    return out, errors


def _check_balanced(balance, errors):
    total = round(sum(nd for _, nd in balance.values()), 2)
    if abs(total) >= TOL:
        errors.append(f"La balance n'est pas équilibrée : total débit − total crédit = {_fmt(total)} "
                      "(si le résultat de l'exercice n'est pas inclus, ajoutez les comptes de résultat "
                      "ou le compte de résultat de l'exercice)")


# ---------------------------------------------------------------------- #
# 2. Balance N-1 (à-nouveaux)
# ---------------------------------------------------------------------- #
def prepare_opening_balance(db, path, opening_date):
    balance, errors = _read_balance(db, path)
    warnings = []
    if not balance and not errors:
        errors.append("Aucune ligne exploitable dans le fichier.")
    _check_balanced(balance, errors)

    bilan, pl_net = [], 0.0
    for acc, nd in balance.values():
        if abs(nd) < 1e-9:
            continue
        if effective_section(acc) in INCOME_SECTIONS + EXPENSE_SECTIONS:
            pl_net += nd
        else:
            bilan.append((acc, nd))
    pl_net = round(pl_net, 2)
    re_acc = None
    if abs(pl_net) >= TOL:
        re_acc = _Resolver(db).by_number.get("1920")
        if not re_acc:
            errors.append("Le fichier contient des comptes de résultat, mais le compte 1920 "
                          "(Retained earnings) est introuvable pour y reporter le résultat N-1.")
        else:
            bilan.append((re_acc, pl_net))
            warnings.append(f"Les comptes de résultat N-1 (résultat {_fmt(-pl_net)}) sont reportés "
                            f"sur « {re_acc['name']} » (1920).")

    existing = db.query("SELECT id FROM journal_entries WHERE doc_type=?", (DOC_OPENING,))
    # fusionner les lignes d'un même compte (RE éventuel)
    merged = {}
    for acc, nd in bilan:
        merged[acc["id"]] = (acc, round(merged.get(acc["id"], (acc, 0.0))[1] + nd, 2))
    lines = [(a, nd) for a, nd in merged.values() if abs(nd) >= 1e-9]
    td = round(sum(nd for _, nd in lines if nd > 0), 2)
    summary = (f"Balance N-1 : {len(lines)} compte(s) au bilan, à-nouveaux au {opening_date}\n"
               f"Total débit = total crédit = {_fmt(td)}")
    if existing:
        warnings.append("Une balance N-1 a déjà été importée : elle sera REMPLACÉE.")

    def apply():
        cn = db.conn
        try:
            for e in existing:
                cn.execute("DELETE FROM journal_lines WHERE entry_id=?", (e["id"],))
                cn.execute("DELETE FROM journal_entries WHERE id=?", (e["id"],))
            cur = cn.execute(
                "INSERT INTO journal_entries (date, doc_type, doc_number, memo) VALUES (?,?,?,?)",
                (opening_date, DOC_OPENING, "N-1", "À-nouveaux — balance N-1"))
            for acc, nd in lines:
                cn.execute(
                    "INSERT INTO journal_lines (entry_id, account_id, description, debit, credit) "
                    "VALUES (?,?,?,?,?)",
                    (cur.lastrowid, acc["id"], "À-nouveaux N-1",
                     nd if nd > 0 else 0, -nd if nd < 0 else 0))
            cn.commit()
        except Exception:
            cn.rollback()
            raise
        return f"Balance N-1 importée ({len(lines)} comptes)."

    return Plan(summary, _errors_head(errors) if errors else [], warnings,
                apply if not errors and lines else None, needs_replace=bool(existing))


# ---------------------------------------------------------------------- #
# 3. Balance N (soldes de clôture)
# ---------------------------------------------------------------------- #
def prepare_closing_balance(db, path, closing_date, zero_missing=True):
    """
    Les soldes du fichier sont des soldes CUMULÉS à la date de clôture. Le logiciel les compare
    à ce qui est déjà dans le journal et passe UNE écriture de régularisation pour l'écart.
    Ré-importer le même fichier ne crée donc rien de plus.
    """
    balance, errors = _read_balance(db, path)
    warnings = []
    if not balance and not errors:
        errors.append("Aucune ligne exploitable dans le fichier.")
    _check_balanced(balance, errors)

    ledger = {r["account_id"]: round(r["nd"], 2) for r in db.query(
        "SELECT account_id, SUM(debit)-SUM(credit) AS nd FROM journal_lines GROUP BY account_id")}
    accounts = {a["id"]: a for a in db.query("SELECT * FROM accounts")}
    diffs = []
    for acc_id, acc in accounts.items():
        in_file = acc_id in balance
        if not in_file and not zero_missing:
            continue
        target = balance[acc_id][1] if in_file else 0.0
        diff = round(target - ledger.get(acc_id, 0.0), 2)
        if abs(diff) >= TOL:
            diffs.append((acc, diff))
    if not zero_missing:
        total = round(sum(d for _, d in diffs), 2)
        if abs(total) >= TOL:
            errors.append(f"La régularisation ne s'équilibre pas (écart {_fmt(total)}) : "
                          "cochez « remettre à zéro les comptes absents du fichier ».")
    missing = [a for a_id, a in accounts.items()
               if a_id not in balance and abs(ledger.get(a_id, 0.0)) >= TOL]
    if zero_missing and missing:
        warnings.append(f"{len(missing)} compte(s) ont un solde en base mais ne figurent pas dans le "
                        "fichier : ils seront remis à zéro.")

    td = round(sum(d for _, d in diffs if d > 0), 2)
    if diffs:
        summary = (f"Balance N au {closing_date} : {len(diffs)} compte(s) à régulariser\n"
                   f"Écriture de régularisation : total débit = total crédit = {_fmt(td)}")
    else:
        summary = "Le journal est déjà conforme à cette balance N : rien à importer."

    def apply():
        cn = db.conn
        try:
            cur = cn.execute(
                "INSERT INTO journal_entries (date, doc_type, doc_number, memo) VALUES (?,?,?,?)",
                (closing_date, DOC_CLOSING, "N", "Régularisation — balance N"))
            for acc, d in diffs:
                cn.execute(
                    "INSERT INTO journal_lines (entry_id, account_id, description, debit, credit) "
                    "VALUES (?,?,?,?,?)",
                    (cur.lastrowid, acc["id"], "Régularisation balance N",
                     d if d > 0 else 0, -d if d < 0 else 0))
            cn.commit()
        except Exception:
            cn.rollback()
            raise
        return f"Balance N importée ({len(diffs)} comptes régularisés)."

    return Plan(summary, _errors_head(errors) if errors else [], warnings,
                apply if not errors and diffs else None)


# ---------------------------------------------------------------------- #
# Modèles Excel
# ---------------------------------------------------------------------- #
def export_template(kind, path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    if kind == "entries":
        ws.title = "Ecritures"
        ws.append(["Date", "N° pièce", "Libellé", "Compte", "Débit", "Crédit"])
        ws.append(["2026-01-05", "OD-001", "Apport initial", "1100", 100000, 0])
        ws.append(["2026-01-05", "OD-001", "Apport initial", "1900", 0, 100000])
        widths = (12, 12, 40, 12, 14, 14)
    else:
        ws.title = "Balance"
        ws.append(["Compte", "Intitulé", "Débit", "Crédit"])
        ws.append(["1100", "Cash in bank", 100000, 0])
        ws.append(["1900", "Common stock", 0, 100000])
        widths = (14, 40, 14, 14)
    for col, w in zip("ABCDEF", widths):
        ws.column_dimensions[col].width = w
    wb.save(path)
